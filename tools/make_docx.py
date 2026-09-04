"""Convert docs/design-intent.md into dist/design-intent.docx.

Needs python-docx (pip install python-docx). If it is not installed this exits
with a message and the submission notes tell the human to paste the seven
sections into the official template by hand instead.

Usage: python tools/make_docx.py
"""

import os
import re
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "docs", "design-intent.md")
DIST = os.path.join(ROOT, "dist")
OUT = os.path.join(DIST, "design-intent.docx")


def parse(text):
    """Return [(heading, [paragraphs])] for each '## ' section, in order."""
    sections = []
    current = None
    for line in text.splitlines():
        if line.startswith("## "):
            current = (line[3:].strip(), [])
            sections.append(current)
        elif line.startswith("# "):
            continue
        elif line.strip() and current is not None:
            current[1].append(line.strip())
    return sections


def main():
    if not os.path.isfile(SRC):
        print("docs/design-intent.md is missing")
        return 1

    text = open(SRC, encoding="utf-8").read()
    sections = parse(text)
    if len(sections) != 7:
        print("expected 7 sections, found %d" % len(sections))
        return 1

    body_words = len(re.sub(r"^#.*$", "", text, flags=re.M).split())
    total_words = len(text.split())
    print("design intent: %d body words, %d including headings" % (body_words, total_words))
    if total_words > 500:
        print("FAIL: over the 500 word limit")
        return 1

    try:
        from docx import Document
    except ImportError:
        print("python-docx is not installed")
        return 2

    os.makedirs(DIST, exist_ok=True)
    doc = Document()
    doc.add_heading("Design Intent", level=0)
    for heading, paragraphs in sections:
        doc.add_heading(heading, level=1)
        for p in paragraphs:
            doc.add_paragraph(p)
    doc.save(OUT)
    print("wrote dist/design-intent.docx")
    return 0


if __name__ == "__main__":
    sys.exit(main())
